"""Convert the author's Cartimar workbook into observation CSV."""

import argparse
import csv
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from .store import fmt


FIELDS = ("date", "sheet", "item_ja", "name_tl", "name_en", "price", "market")


def _date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return date(1899, 12, 30) + timedelta(days=value)
    return datetime.strptime(str(value).strip(), "%m/%d/%Y").date()


def read_own(xlsx: Path) -> list[dict]:
    workbook = load_workbook(xlsx, data_only=True, read_only=True)
    rows = []
    try:
        for sheet in workbook:
            if sheet.title not in ("野菜", "魚介"):
                continue
            values = sheet.iter_rows(values_only=True)
            header = next(values)
            dates = [_date(value) for value in header[3:]]
            for cells in values:
                item = str(cells[2]).strip() if cells[2] is not None else ""
                if not item:
                    continue
                for day, price in zip(dates, cells[3:]):
                    if not isinstance(price, (int, float)) or isinstance(price, bool):
                        continue
                    rows.append(dict(date=day.isoformat(), sheet=sheet.title, item_ja=item,
                                     name_tl=str(cells[1]).strip() if cells[1] is not None else "",
                                     name_en=str(cells[0]).strip() if cells[0] is not None else "",
                                     price=float(price), market="Cartimar"))
    finally:
        workbook.close()
    return rows


def main(argv=None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--out", type=Path, default=Path("data/own.csv"))
    args = parser.parse_args(argv)
    rows = sorted(read_own(args.xlsx), key=lambda row: (row["sheet"], row["item_ja"], row["date"]))
    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows({**row, "price": fmt(row["price"])} for row in rows)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
